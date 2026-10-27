import base64
import hashlib
import hmac
import io
import json
import logging
import os
import re
import secrets
import sqlite3
import time
from collections import Counter
from pathlib import Path
from typing import Optional

from dotenv import load_dotenv
from fastapi import Depends, FastAPI, File, Form, Header, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response, StreamingResponse
from groq import Groq
from pydantic import BaseModel, EmailStr
import pdfplumber
from docx import Document as DocxDocument
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

load_dotenv()
app = FastAPI(title="Resume Customizer API")
logger = logging.getLogger(__name__)
app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:3000",
        "http://127.0.0.1:3000",
        "https://ai-resume-customizer-cywj.vercel.app",
    ],
    allow_origin_regex=r"https://ai-resume-customizer-cywj-[a-z0-9-]+-team-c25\.vercel\.app",
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

MAX_FILE = 5 * 1024 * 1024
DB_PATH = Path(os.getenv("DATABASE_PATH", Path(__file__).with_name("resume_customizer.sqlite3")))
TOKEN_SECRET = os.getenv("AUTH_SECRET")
VALID_SECTIONS = {"summary", "skills", "experience", "education", "projects"}


def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    with get_db() as conn:
        conn.execute("""CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            email TEXT UNIQUE NOT NULL,
            password_hash TEXT NOT NULL,
            created_at INTEGER NOT NULL
        )""")
        conn.execute("""CREATE TABLE IF NOT EXISTS history (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
            title TEXT NOT NULL,
            kind TEXT NOT NULL,
            job_description TEXT NOT NULL,
            original_resume TEXT NOT NULL,
            result TEXT NOT NULL,
            match_score INTEGER NOT NULL,
            ats_score INTEGER NOT NULL,
            created_at INTEGER NOT NULL
        )""")
        conn.execute("CREATE TABLE IF NOT EXISTS app_config (key TEXT PRIMARY KEY, value TEXT NOT NULL)")
        conn.execute("PRAGMA foreign_keys = ON")
        conn.execute("INSERT OR IGNORE INTO app_config (key, value) VALUES ('token_secret', ?)", (secrets.token_urlsafe(48),))


init_db()
if not TOKEN_SECRET:
    with get_db() as conn:
        TOKEN_SECRET = conn.execute("SELECT value FROM app_config WHERE key = 'token_secret'").fetchone()[0]


class Credentials(BaseModel):
    email: EmailStr
    password: str


def password_digest(password: str, salt: Optional[bytes] = None) -> str:
    salt = salt or secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode(), salt, 240_000)
    return f"{base64.urlsafe_b64encode(salt).decode()}:{base64.urlsafe_b64encode(digest).decode()}"


def check_password(password: str, stored: str) -> bool:
    try:
        salt_text, digest_text = stored.split(":", 1)
        salt = base64.urlsafe_b64decode(salt_text)
        expected = base64.urlsafe_b64decode(digest_text)
        actual = hashlib.pbkdf2_hmac("sha256", password.encode(), salt, 240_000)
        return hmac.compare_digest(actual, expected)
    except (ValueError, TypeError):
        return False


def issue_token(user_id: int, email: str) -> str:
    payload = base64.urlsafe_b64encode(json.dumps({"sub": user_id, "email": email, "exp": int(time.time()) + 60 * 60 * 24 * 30}).encode()).decode().rstrip("=")
    signature = hmac.new(TOKEN_SECRET.encode(), payload.encode(), hashlib.sha256).hexdigest()
    return f"{payload}.{signature}"


def get_current_user(authorization: Optional[str] = Header(default=None)):
    if not authorization or not authorization.startswith("Bearer "):
        return None
    try:
        payload, signature = authorization[7:].split(".", 1)
        expected = hmac.new(TOKEN_SECRET.encode(), payload.encode(), hashlib.sha256).hexdigest()
        if not hmac.compare_digest(signature, expected):
            return None
        data = json.loads(base64.urlsafe_b64decode(payload + "=" * (-len(payload) % 4)))
        if data.get("exp", 0) < time.time():
            return None
        with get_db() as conn:
            row = conn.execute("SELECT id, email FROM users WHERE id = ?", (data["sub"],)).fetchone()
        return dict(row) if row else None
    except (ValueError, KeyError, TypeError, json.JSONDecodeError):
        return None


def require_user(user=Depends(get_current_user)):
    if not user:
        raise HTTPException(401, "Sign in to use this feature.")
    return user


@app.post("/auth/register")
def register(credentials: Credentials):
    if len(credentials.password) < 8:
        raise HTTPException(400, "Password must be at least 8 characters.")
    email = str(credentials.email).lower().strip()
    try:
        with get_db() as conn:
            cursor = conn.execute("INSERT INTO users (email, password_hash, created_at) VALUES (?, ?, ?)", (email, password_digest(credentials.password), int(time.time())))
            user_id = cursor.lastrowid
    except sqlite3.IntegrityError:
        raise HTTPException(409, "An account with this email already exists.")
    return {"token": issue_token(user_id, email), "user": {"id": user_id, "email": email}}


@app.post("/auth/login")
def login(credentials: Credentials):
    email = str(credentials.email).lower().strip()
    with get_db() as conn:
        user = conn.execute("SELECT id, email, password_hash FROM users WHERE email = ?", (email,)).fetchone()
    if not user or not check_password(credentials.password, user["password_hash"]):
        raise HTTPException(401, "Email or password is incorrect.")
    return {"token": issue_token(user["id"], user["email"]), "user": {"id": user["id"], "email": user["email"]}}


@app.get("/auth/me")
def me(user=Depends(require_user)):
    return {"user": user}


def parse_pdf(content: bytes) -> str:
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


def parse_docx(content: bytes) -> str:
    doc = DocxDocument(io.BytesIO(content))
    parts = []
    for child in doc.element.body:
        if child.tag.endswith("}p"):
            parts.append(DocxParagraph(child, doc).text)
        elif child.tag.endswith("}tbl"):
            table = DocxTable(child, doc)
            parts.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
    for section in doc.sections:
        parts.extend(paragraph.text for paragraph in section.header.paragraphs)
        parts.extend(paragraph.text for paragraph in section.footer.paragraphs)
    return "\n".join(part for part in parts if part.strip())


def extract_resume(file: UploadFile, content: bytes) -> str:
    filename = file.filename or ""
    if filename.lower().endswith(".pdf"):
        resume = parse_pdf(content)
    elif filename.lower().endswith(".docx"):
        resume = parse_docx(content)
    else:
        raise HTTPException(400, "Unsupported file type (PDF/DOCX only).")
    if not resume.strip():
        raise HTTPException(400, "Could not extract text from resume.")
    return resume


STOP_WORDS = set("a an and are as at be by for from has have in into is it of on or that the their this to with your you our we will role work working experience skills ability strong using including required preferred candidate team".split())


def terms(text: str):
    return [word for word in re.findall(r"[a-zA-Z][a-zA-Z+#.\-/]{1,}", text.lower()) if word not in STOP_WORDS]


def calculate_match(resume: str, jd: str) -> int:
    resume_terms, jd_terms = set(terms(resume)), set(terms(jd))
    if not jd_terms:
        return 0
    return min(100, round(100 * len(resume_terms & jd_terms) / len(jd_terms)))


def calculate_ats(text: str) -> int:
    score = 35
    lower = text.lower()
    if len(text) > 700:
        score += 15
    if len(text) > 1400:
        score += 10
    headings = ("summary", "skills", "experience", "education", "projects")
    score += min(25, sum(5 for heading in headings if re.search(rf"\b{heading}\b", lower)))
    score += min(10, len(re.findall(r"\b\d+(?:\.\d+)?%?\b", text)) * 2)
    if "|" not in text and "●" not in text and "★" not in text:
        score += 5
    return max(0, min(100, score))


def save_history(user, title, kind, jd, original, result, match_score, ats_score):
    if not user:
        return None
    with get_db() as conn:
        cursor = conn.execute(
            "INSERT INTO history (user_id, title, kind, job_description, original_resume, result, match_score, ats_score, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (user["id"], title[:120], kind, jd, original, result, match_score, ats_score, int(time.time())),
        )
        return cursor.lastrowid


def get_client():
    key = os.getenv("GROQ_API_KEY")
    if not key:
        raise HTTPException(500, "GROQ_API_KEY is not configured on the backend.")
    return Groq(api_key=key)


def stream_completion(client, prompt: str, temperature: float):
    model = os.getenv("GROQ_MODEL", "llama-3.3-70b-versatile")
    options = {
        "model": model,
        "messages": [{"role": "user", "content": prompt}],
        "temperature": temperature,
        "stream": True,
    }
    # Keep reasoning models from spending their entire completion budget on
    # hidden reasoning before producing the user-visible answer.
    if model.startswith("openai/gpt-oss"):
        options["reasoning_effort"] = "low"
        options["max_completion_tokens"] = 8192
    return client.chat.completions.create(**options)


def text_from_content(content) -> str:
    if isinstance(content, str):
        return content
    if isinstance(content, list):
        return "".join(
            part.get("text", "") if isinstance(part, dict) else getattr(part, "text", "")
            for part in content
        )
    return ""


def make_prompt(jd: str, resume: str, selected: set[str], kind: str) -> str:
    if kind == "cover_letter":
        return f"""Write a polished, specific, one-page cover letter using only truthful information in the candidate resume. Do not invent experience, employers, dates, credentials, or metrics. Address the requirements in the job description and use a professional, human tone. Return only the cover letter.\n\nJOB DESCRIPTION:\n{jd}\n\nCANDIDATE RESUME:\n{resume}"""
    selected_text = ", ".join(sorted(selected))
    return f"""You are an expert resume writer and ATS optimizer. Produce a complete tailored resume in plain text. Customize ONLY these sections: {selected_text}. Preserve all non-selected sections and their facts as faithfully as possible. Reorder content only where helpful. Never invent or exaggerate employers, dates, credentials, tools, responsibilities, or metrics. Keep measurable facts from the source. Use standard ATS-friendly section headings, simple bullets, and no columns/tables/icons. Do not add commentary or a suggested keywords section.\n\nJOB DESCRIPTION:\n{jd}\n\nORIGINAL RESUME:\n{resume}"""


@app.post("/customize/stream")
async def customize_stream(
    jd: str = Form(...),
    file: UploadFile = File(...),
    sections: str = Form("summary,skills,experience,education,projects"),
    title: str = Form("Tailored resume"),
    authorization: Optional[str] = Header(default=None),
):
    if len(jd.strip()) < 20:
        raise HTTPException(400, "Job description must be at least 20 characters.")
    content = await file.read()
    if len(content) > MAX_FILE:
        raise HTTPException(400, "File too large (max 5MB).")
    resume = extract_resume(file, content)
    selected = {part.strip().lower() for part in sections.split(",")} & VALID_SECTIONS
    if not selected:
        raise HTTPException(400, "Select at least one section to customize.")
    user = get_current_user(authorization)
    original_match = calculate_match(resume, jd)
    client = get_client()
    prompt = make_prompt(jd, resume, selected, "resume")

    def events():
        result = ""
        finish_reason = None
        try:
            stream = stream_completion(client, prompt, 0.3)
            for chunk in stream:
                if not chunk.choices:
                    continue
                choice = chunk.choices[0]
                finish_reason = choice.finish_reason or finish_reason
                delta = text_from_content(choice.delta.content)
                if delta:
                    result += delta
                    yield f"event: chunk\ndata: {json.dumps({'text': delta})}\n\n"
            if not result.strip():
                model = os.getenv("GROQ_MODEL", "llama-3.3-70b-versatile")
                logger.warning("LLM returned no visible resume text (model=%s, finish_reason=%s)", model, finish_reason)
                yield f"event: error\ndata: {json.dumps({'message': f'The model {model} finished without returning resume text (finish reason: {finish_reason or "unknown"}). Try lowering reasoning effort or selecting a different GROQ_MODEL.'})}\n\n"
                return
            tailored_match = calculate_match(result, jd)
            ats_score = calculate_ats(result)
            history_id = save_history(user, title, "resume", jd, resume, result, tailored_match, ats_score)
            metadata = {"original": resume, "filename": file.filename, "originalMatchScore": original_match, "matchScore": tailored_match, "atsScore": ats_score, "historyId": history_id}
            yield f"event: done\ndata: {json.dumps(metadata)}\n\n"
        except Exception as exc:
            logger.exception("Resume generation stream failed")
            message = str(exc)[:300] if isinstance(exc, HTTPException) else f"Resume generation failed ({type(exc).__name__}). Check the configured GROQ_MODEL and API key, then try again."
            yield f"event: error\ndata: {json.dumps({'message': message})}\n\n"

    return StreamingResponse(events(), media_type="text/event-stream", headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


@app.post("/cover-letter/stream")
async def cover_letter_stream(
    jd: str = Form(...),
    file: UploadFile = File(...),
    title: str = Form("Cover letter"),
    authorization: Optional[str] = Header(default=None),
):
    if len(jd.strip()) < 20:
        raise HTTPException(400, "Job description must be at least 20 characters.")
    content = await file.read()
    if len(content) > MAX_FILE:
        raise HTTPException(400, "File too large (max 5MB).")
    resume = extract_resume(file, content)
    user = get_current_user(authorization)
    client = get_client()
    prompt = make_prompt(jd, resume, set(), "cover_letter")

    def events():
        result = ""
        finish_reason = None
        try:
            stream = stream_completion(client, prompt, 0.4)
            for chunk in stream:
                if not chunk.choices:
                    continue
                choice = chunk.choices[0]
                finish_reason = choice.finish_reason or finish_reason
                delta = text_from_content(choice.delta.content)
                if delta:
                    result += delta
                    yield f"event: chunk\ndata: {json.dumps({'text': delta})}\n\n"
            if not result.strip():
                model = os.getenv("GROQ_MODEL", "llama-3.3-70b-versatile")
                logger.warning("LLM returned no visible cover letter text (model=%s, finish_reason=%s)", model, finish_reason)
                yield f"event: error\ndata: {json.dumps({'message': f'The model {model} finished without returning cover letter text (finish reason: {finish_reason or "unknown"}). Try lowering reasoning effort or selecting a different GROQ_MODEL.'})}\n\n"
                return
            match_score = calculate_match(resume, jd)
            ats_score = calculate_ats(resume)
            history_id = save_history(user, title, "cover_letter", jd, resume, result, match_score, ats_score)
            yield f"event: done\ndata: {json.dumps({'original': resume, 'filename': file.filename, 'matchScore': match_score, 'atsScore': ats_score, 'historyId': history_id})}\n\n"
        except Exception as exc:
            logger.exception("Cover letter generation stream failed")
            yield f"event: error\ndata: {json.dumps({'message': f'Cover letter generation failed ({type(exc).__name__}). Check the configured GROQ_MODEL and API key, then try again.'})}\n\n"

    return StreamingResponse(events(), media_type="text/event-stream", headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


@app.get("/history")
def history(user=Depends(require_user)):
    with get_db() as conn:
        rows = conn.execute("SELECT id, title, kind, job_description, original_resume, result, match_score, ats_score, created_at FROM history WHERE user_id = ? ORDER BY created_at DESC, id DESC", (user["id"],)).fetchall()
    return {"items": [dict(row) for row in rows]}


@app.delete("/history/{history_id}")
def delete_history(history_id: int, user=Depends(require_user)):
    with get_db() as conn:
        cursor = conn.execute("DELETE FROM history WHERE id = ? AND user_id = ?", (history_id, user["id"]))
    if cursor.rowcount == 0:
        raise HTTPException(404, "History item not found.")
    return {"deleted": True}


@app.post("/download")
async def download(text: str = Form(...), format: str = Form("txt")):
    if format == "txt":
        return Response(content=text, media_type="text/plain", headers={"Content-Disposition": "attachment; filename=resume.txt"})
    if format == "docx":
        doc = DocxDocument()
        for line in text.splitlines():
            doc.add_paragraph(line)
        buffer = io.BytesIO()
        doc.save(buffer)
        return Response(content=buffer.getvalue(), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": "attachment; filename=resume.docx"})
    if format == "pdf":
        buffer = io.BytesIO()
        pdf = canvas.Canvas(buffer, pagesize=letter)
        width, height = letter
        y = height - 48
        for line in text.splitlines():
            chunks = [line[i:i + 95] for i in range(0, max(1, len(line)), 95)]
            for chunk in chunks:
                if y < 48:
                    pdf.showPage()
                    y = height - 48
                pdf.drawString(48, y, chunk)
                y -= 14
        pdf.save()
        return Response(content=buffer.getvalue(), media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=resume.pdf"})
    raise HTTPException(400, "Unsupported format.")

from fastapi.responses import StreamingResponse

@app.post("/customize/stream")
async def customize_stream(jd: str = Form(...), file: UploadFile = File(...)):
    content = await file.read()
    name = file.filename or ""
    lower = name.lower()
    if lower.endswith(".pdf"):
        resume = parse_pdf(content)
    elif lower.endswith(".docx"):
        resume = parse_docx(content)
    else:
        raise HTTPException(400, "Unsupported file type")
    key = os.getenv("GROQ_API_KEY")
    client = Groq(api_key=key)
    prompt = f"Tailor resume to JD. JD: {jd} Resume: {resume}"
    def gen():
        r = client.chat.completions.create(model=os.getenv("GROQ_MODEL","llama-3.3-70b-versatile"), messages=[{"role":"user","content":prompt}], temperature=0.4, stream=True)
        for chunk in r:
            delta = chunk.choices[0].delta.content or ""
            yield delta
    return StreamingResponse(gen(), media_type="text/plain")
